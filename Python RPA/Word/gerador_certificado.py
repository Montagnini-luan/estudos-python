from docx import Document
from docx.shared import Pt
from openpyxl import load_workbook
import os

nome_arquivo_alunos = "C:\\Users\\luanm\\OneDrive\\Desktop\\Sites\\Estudos\\estudos python\\Python RPA\\Word\\Alunos.xlsx"

planilha_dados_alunos = load_workbook(nome_arquivo_alunos)

sheet_selecionada = planilha_dados_alunos['Nomes']

for linha in range(2, len(sheet_selecionada["A"]) + 1):

    arquivo_word = Document("C:\\Users\\luanm\\OneDrive\\Desktop\\Sites\\Estudos\\estudos python\\Python RPA\\Word\\Certificado1.docx")

    estilo = arquivo_word.styles["Normal"]

    nome_aluno = sheet_selecionada['A%s' % linha].value

    for paragrafo in arquivo_word.paragraphs:

        if "@nome" in paragrafo.text:

            paragrafo.text = nome_aluno
            fonte = estilo.font
            fonte.name = "Calibri (Corpo)"
            fonte.size = Pt(24)

    caminho_certificados = "C:\\Users\\luanm\OneDrive\\Desktop\\Sites\\Estudos\\estudos python\\Python RPA\\Word\\Certificados\\" + nome_aluno + ".docx"

    arquivo_word.save(caminho_certificados)

print("Certificados gerados com sucesso")
